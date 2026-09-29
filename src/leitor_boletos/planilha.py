import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .linha import (MOEDA_NAO_SUPORTADA, VALOR_DE_REFERENCIA, VALOR_E_VENCIMENTO_988,
                    VALOR_EM_ABERTO, Boleto, formatar)
from .tabela import DUPLICADO, INVALIDO, Tabela

COLUNAS = ["Arquivo", "Banco", "Valor", "Vencimento", "Situação", "Linha digitável", "Aviso"]
LARGURAS = [26, 36, 16, 19, 19, 60, 42]
FORMATO_REAIS = '"R$" #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"


def nome_sugerido(hoje: date, extensao: str) -> str:
    return f"boletos-{hoje:%Y-%m-%d}.{extensao}"


def aviso_texto(boleto: Boleto) -> str:
    avisos = list(boleto.avisos)
    if boleto.valor is None and boleto.motivo and boleto.motivo not in avisos:
        avisos.insert(0, boleto.motivo)
    return "; ".join(avisos)


NOMES_CURTOS = {
    INVALIDO: "com problema",
    DUPLICADO: "duplicado",
    VALOR_EM_ABERTO: "valor em aberto",
    VALOR_DE_REFERENCIA: "valor de referência",
    VALOR_E_VENCIMENTO_988: "código 988",
    MOEDA_NAO_SUPORTADA: "moeda não suportada",
}


def descrever_fora(tabela: Tabela) -> str:
    return ", ".join(f"{qtd} {NOMES_CURTOS.get(motivo, motivo)}"
                     for motivo, qtd in tabela.fora_do_total.items())


def _linha_digitavel(boleto: Boleto) -> str:
    return formatar(boleto.numero) if boleto.numero else ""


def exportar_xlsx(tabela: Tabela, caminho: Path | str) -> None:
    livro = Workbook()
    folha = livro.active
    folha.title = "Boletos"
    folha.append(COLUNAS)
    for celula in folha[1]:
        celula.font = Font(bold=True)

    for boleto in tabela.linhas:
        folha.append([
            boleto.origem,
            boleto.banco or "",
            boleto.valor,
            boleto.vencimento if boleto.vencimento else boleto.vencimento_texto,
            boleto.situacao,
            _linha_digitavel(boleto),
            aviso_texto(boleto),
        ])
        linha = folha.max_row
        folha.cell(linha, 3).number_format = FORMATO_REAIS
        if boleto.vencimento:
            folha.cell(linha, 4).number_format = FORMATO_DATA

    ultima = folha.max_row
    folha.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{ultima}"
    folha.freeze_panes = "A2"
    for indice, largura in enumerate(LARGURAS, start=1):
        folha.column_dimensions[get_column_letter(indice)].width = largura

    folha.append([])
    folha.append(["Total", "", tabela.total])
    total = folha.max_row
    folha.cell(total, 1).font = Font(bold=True)
    folha.cell(total, 3).font = Font(bold=True)
    folha.cell(total, 3).number_format = FORMATO_REAIS
    if tabela.fora_do_total:
        folha.append(["Fora do total", descrever_fora(tabela)])

    livro.save(caminho)


def _valor_csv(boleto: Boleto) -> str:
    return "" if boleto.valor is None else f"{boleto.valor:.2f}".replace(".", ",")


def exportar_csv(tabela: Tabela, caminho: Path | str) -> None:
    with open(caminho, "w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(COLUNAS)
        for boleto in tabela.linhas:
            escritor.writerow([
                boleto.origem,
                boleto.banco or "",
                _valor_csv(boleto),
                boleto.vencimento_texto,
                boleto.situacao,
                _linha_digitavel(boleto),
                aviso_texto(boleto),
            ])
