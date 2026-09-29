import csv
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from leitor_boletos import linha as L
from leitor_boletos.pdf import ler_pdf
from leitor_boletos.planilha import COLUNAS, exportar_csv, exportar_xlsx, nome_sugerido
from leitor_boletos.tabela import montar

HOJE = date(2026, 9, 29)
EXEMPLOS = Path(__file__).resolve().parent.parent / "exemplos"


def tabela_dos_exemplos():
    return montar([b for arquivo in sorted(EXEMPLOS.glob("*.pdf")) for b in ler_pdf(arquivo, HOJE)])


def test_xlsx_colunas_tipos_e_total(tmp_path):
    destino = tmp_path / "boletos.xlsx"
    exportar_xlsx(tabela_dos_exemplos(), destino)
    folha = load_workbook(destino).active

    assert [c.value for c in folha[1]] == COLUNAS
    assert folha["A2"].value == "01-banco-normal.pdf"
    assert folha["C2"].value == 350
    assert folha["C2"].number_format == '"R$" #,##0.00'
    assert folha["D2"].value == datetime(2026, 10, 15)
    assert folha["D2"].number_format == "DD/MM/YYYY"
    assert folha.auto_filter.ref == "A1:G15"

    valor_em_aberto = folha["C6"]
    assert valor_em_aberto.value is None
    assert folha["G6"].value == L.VALOR_EM_ABERTO

    total = [linha for linha in folha.iter_rows(values_only=True) if linha[0] == "Total"]
    assert total and Decimal(str(total[0][2])) == Decimal("2457.00")


def test_csv_padrao_do_excel_brasileiro(tmp_path):
    destino = tmp_path / "boletos.csv"
    exportar_csv(tabela_dos_exemplos(), destino)
    bruto = destino.read_bytes()
    assert bruto.startswith(b"\xef\xbb\xbf")

    linhas = list(csv.reader(destino.read_text(encoding="utf-8-sig").splitlines(), delimiter=";"))
    assert linhas[0] == COLUNAS
    assert linhas[2][2] == "1249,90"
    assert linhas[1][3] == "15/10/2026"
    assert len(linhas) == 15


def test_exportar_de_novo_por_cima(tmp_path):
    destino = tmp_path / "boletos.xlsx"
    exportar_xlsx(montar([L.interpretar("10490055057722213334877777777713432420000032112", HOJE)]), destino)
    exportar_xlsx(tabela_dos_exemplos(), destino)
    assert load_workbook(destino).active.max_row > 15


def test_nome_sugerido():
    assert nome_sugerido(HOJE, "xlsx") == "boletos-2026-09-29.xlsx"
